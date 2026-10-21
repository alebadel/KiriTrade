#!/usr/bin/env python3
"""Build data/corpus.json: an inventory of the documents the tracker's AI agents can read.

Counts only. No document text is written to the output, so the public site never
carries the content of a copyrighted paper or a working draft. Each row carries a
title (public documents have public titles), a document type, the country the
document is about (ISO3, or GEN for general material), the year, pages when known,
and the number of words in the text the knowledge base holds.

Sources read (all inside the KiriTrade folder):
  standards_docs/_KB/INDEX.md and standards_docs/_KB/md/*.md   (standards, customs rules, data documentation)
  mirror_papers/_KB/INDEX_ALL.md                               (academic papers and IMF reports on mirror data)
  *.docx and *.md in the KiriTrade root listed in PROJECT_NOTES (the author's own briefs)
  country_docs/<ISO3>/**                                       (optional: any documents dropped there, by country)

Usage:  python tracker/build_corpus_data.py --root <KiriTrade folder> --out data/corpus.json
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

TYPES = {
    "standards": "Statistical manuals and standards",
    "customs": "Customs rules and regulations",
    "datadoc": "Data documentation",
    "academic": "Academic papers",
    "imf": "IMF and other official reports",
    "notes": "Project briefs and drafts",
    "news": "News",
}
TYPE_ORDER = ["standards", "customs", "datadoc", "academic", "imf", "notes", "news"]

COUNTRY_NAMES = {
    "GEN": "General (all countries)", "KIR": "Kiribati", "THA": "Thailand", "JPN": "Japan", "KOR": "Korea",
    "PHL": "Philippines", "IDN": "Indonesia", "ECU": "Ecuador", "COL": "Colombia", "CHN": "China",
    "KHM": "Cambodia", "LAO": "Lao P.D.R.", "KGZ": "Kyrgyz Republic", "NPL": "Nepal", "IND": "India",
    "USA": "United States", "VNM": "Viet Nam", "PRK": "North Korea",
}

# The author's own briefs in the KiriTrade root. Titles are given here so the file names stay private.
PROJECT_NOTES = [
    ("Kiribati_mirror_trade_paper_1.docx", "Mirror-trade paper, part 1: what partner data say about Kiribati's trade", "KIR", 2026),
    ("Kiribati_mirror_trade_paper_2_fishing_fees.docx", "Mirror-trade paper, part 2: fishing fees and the tuna series", "KIR", 2026),
    ("Kiribati_two_flows_of_fish.docx", "Two flows of fish: flag-based and residence-based records", "KIR", 2026),
    ("Kiribati_fishing_fees_vs_mirror_tuna.docx", "Fishing fees and partner-reported tuna: why they are different flows", "KIR", 2026),
    ("Kiribati_fish_exports_to_May2026_brief.docx", "Brief: partner-reported fish exports to May 2026", "KIR", 2026),
    ("Kiribati_rice_sugar_imports_brief.docx", "Brief: rice and sugar imports in partner data", "KIR", 2026),
    ("econ_lit_mirror_20261006.md", "Reading notes on the mirror-data literature", "GEN", 2026),
]

WORD_RE = re.compile(r"[A-Za-z0-9À-ɏ฀-๿]+")
MOJIBAKE = re.compile(r"[─-╿±ΦπΓΣσμτ]")


def count_words(text: str) -> int:
    return len(WORD_RE.findall(text))


def file_date(p: Path) -> str:
    """Date the text entered the knowledge base: the file's modification date (ISO, UTC)."""
    return dt.datetime.fromtimestamp(p.stat().st_mtime, dt.timezone.utc).strftime("%Y-%m-%d")


def md_table_rows(path: Path) -> list[list[str]]:
    rows = []
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if line.startswith("|") and not line.startswith("|---") and not line.startswith("| #"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            rows.append(cells)
    return rows


def link_text(cell: str) -> tuple[str, str]:
    m = re.match(r"\[(.*?)\]\(<?(.*?)>?\)", cell)
    return (m.group(1), m.group(2)) if m else (cell, "")


def standards_rows(root: Path) -> list[dict]:
    kb = root / "standards_docs" / "_KB"
    index = kb / "INDEX.md"
    if not index.exists():
        return []
    out = []
    attach_counter: dict[str, int] = {}
    for cells in md_table_rows(index):
        if len(cells) < 10:
            continue
        n, doc, source, dtype, year, lang, pages, paras, extraction, file = cells[:10]
        title, mdlink = link_text(doc)
        if "unrelated" in title.lower() or "stray" in source.lower():
            continue  # the index itself marks the file as a mistaken download
        file = file.strip("`")
        parts = re.split(r"[\\/]+", file)
        folder = parts[0]
        if folder == "thailand_customs":
            kind, country = "customs", "THA"
            if MOJIBAKE.search(title) or not title.isascii() or re.fullmatch(r"[\d_\-]+(_DFQF)?", title):
                # Thai-language attachment extracted from a notification archive: give it a plain English label
                arch = next((p for p in parts if "Notification" in p), "")
                who = "Ministry of Finance" if "MoF" in arch else "Customs Department"
                m = re.search(r"Notification_(\d{4})-(\d{2})-(\d{2})", arch)
                if m:
                    d = dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
                    ref = "dated " + d.strftime("%d %B %Y").lstrip("0")
                else:
                    m = re.search(r"Notification_(\d+-\d+)", arch)
                    ref = "No. " + m.group(1).replace("-", "/") if m else ""
                attach_counter[arch] = attach_counter.get(arch, 0) + 1
                title = f"Thailand {who} Notification {ref}: attachment {attach_counter[arch]} (Thai language)".replace("  ", " ")
                source = "Thailand " + who
                kind = "customs"
        elif folder == "comtrade_metadata":
            kind = "datadoc"
            m = re.search(r"country_notes_([A-Z]{3})", file)
            country = m.group(1) if m else "GEN"
        else:
            kind, country = "standards", "GEN"
        words = None
        added = None
        mdfile = kb / mdlink if mdlink else None
        if mdfile and mdfile.exists():
            words = count_words(mdfile.read_text(encoding="utf-8", errors="replace"))
            added = file_date(mdfile)
        out.append({
            "id": f"std-{int(n):02d}" if n.isdigit() else f"std-{n}",
            "title": title, "type": kind, "country": country,
            "year": int(year) if year.isdigit() else None,
            "pages": int(pages) if pages.isdigit() else None,
            "words": words, "source": source, "language": lang, "added": added,
            "collection": "standards_docs",
        })
    return out


PAPER_COUNTRY = [
    ("Cambodia", "KHM"), ("Lao", "LAO"), ("Kyrgyz", "KGZ"), ("Nepal", "NPL"), ("India", "IND"),
    ("North Korea", "PRK"), ("China", "CHN"), ("Viet", "VNM"),
]


def papers_rows(root: Path) -> list[dict]:
    index = root / "mirror_papers" / "_KB" / "INDEX_ALL.md"
    if not index.exists():
        return []
    out = []
    for cells in md_table_rows(index):
        if len(cells) < 7:
            continue
        n, paper, authors, source, year, pages, words = cells[:7]
        title, mdlink = link_text(paper)
        is_imf = "IMF" in source or mdlink.startswith("md/IMF_")
        kind = "imf" if is_imf else "academic"
        country = "GEN"
        probe = mdlink + " " + title
        for key, iso in PAPER_COUNTRY:
            if key in probe:
                country = iso
                break
        mdfile = root / "mirror_papers" / "_KB" / mdlink if mdlink else None
        out.append({
            "added": file_date(mdfile) if mdfile and mdfile.exists() else None,
            "id": f"pap-{int(n):02d}" if n.isdigit() else f"pap-{n}",
            "title": title.replace("*", "").replace("•", "'").replace("✩", "").strip(), "type": kind, "country": country,
            "year": int(year) if year.isdigit() else None,
            "pages": int(pages) if pages.isdigit() else None,
            "words": int(words) if words.isdigit() else None,
            "source": source, "authors": authors, "collection": "mirror_papers",
        })
    return out


def docx_words(path: Path) -> int | None:
    try:
        import docx  # python-docx
    except ImportError:
        return None
    try:
        d = docx.Document(str(path))
    except Exception:
        return None
    text = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                text += "\n" + c.text
    return count_words(text)


def notes_rows(root: Path) -> list[dict]:
    out = []
    for i, (fname, title, country, year) in enumerate(PROJECT_NOTES, 1):
        p = root / fname
        if not p.exists():
            continue
        words = docx_words(p) if p.suffix.lower() == ".docx" else count_words(p.read_text(encoding="utf-8", errors="replace"))
        out.append({"id": f"note-{i:02d}", "title": title, "type": "notes", "country": country,
                    "year": year, "pages": None, "words": words, "source": "Author's working files",
                    "added": file_date(p), "collection": "project_notes"})
    return out


def country_docs_rows(root: Path) -> list[dict]:
    """Optional folder country_docs/<ISO3>/<type>/file: type folder is one of the TYPES keys."""
    base = root / "country_docs"
    if not base.exists():
        return []
    out = []
    k = 0
    for iso_dir in sorted(base.iterdir()):
        if not iso_dir.is_dir() or len(iso_dir.name) != 3:
            continue
        for p in sorted(iso_dir.rglob("*")):
            if not p.is_file() or p.name.startswith("_"):
                continue
            kind = p.parent.name if p.parent.name in TYPES else "notes"
            words = None
            if p.suffix.lower() == ".docx":
                words = docx_words(p)
            elif p.suffix.lower() in (".md", ".txt", ".html"):
                words = count_words(p.read_text(encoding="utf-8", errors="replace"))
            elif p.suffix.lower() == ".pdf":
                try:
                    from pypdf import PdfReader
                    r = PdfReader(str(p))
                    words = sum(count_words(pg.extract_text() or "") for pg in r.pages)
                except Exception:
                    words = None
            k += 1
            out.append({"id": f"cd-{k:03d}", "title": p.stem.replace("_", " "), "type": kind,
                        "country": iso_dir.name.upper(), "year": None, "pages": None, "words": words,
                        "source": "country_docs folder", "added": file_date(p), "collection": "country_docs"})
    return out


def build(root: Path) -> dict:
    docs = standards_rows(root) + papers_rows(root) + notes_rows(root) + country_docs_rows(root)
    countries = sorted({d["country"] for d in docs}, key=lambda c: (c != "GEN", c != "KIR", c))
    by = {}
    for c in countries:
        by[c] = {k: {"documents": 0, "words": 0, "pages": 0} for k in TYPE_ORDER}
    for d in docs:
        cell = by[d["country"]][d["type"]]
        cell["documents"] += 1
        cell["words"] += d["words"] or 0
        cell["pages"] += d["pages"] or 0
    totals = {k: {"documents": 0, "words": 0} for k in TYPE_ORDER}
    for d in docs:
        totals[d["type"]]["documents"] += 1
        totals[d["type"]]["words"] += d["words"] or 0
    dates = sorted({d["added"] for d in docs if d.get("added")})
    growth = []
    for day in dates:
        row = {"date": day}
        for k in TYPE_ORDER:
            row[k] = sum(d["words"] or 0 for d in docs if d["type"] == k and d.get("added") and d["added"] <= day)
            row[k + "_docs"] = sum(1 for d in docs if d["type"] == k and d.get("added") and d["added"] <= day)
        growth.append(row)
    return {
        "growth": growth,
        "meta": {
            "generated_at": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
            "builder": "tracker/build_corpus_data.py",
            "what": "Inventory of the documents available to the tracker's AI agents: counts of documents and words by type and by the country the document is about. No document text is published.",
            "types": TYPES, "type_order": TYPE_ORDER, "country_names": COUNTRY_NAMES,
            "n_documents": len(docs), "n_words": sum(d["words"] or 0 for d in docs),
            "notes": [
                "added = the date the document's text entered the knowledge base (modification date of the converted text file); the growth chart on the site is cumulative by this date.",
                "Words are counted in the text the knowledge base holds (converted from PDF, HTML or Word files); scanned pages with no text layer count as zero.",
                "GEN marks material that is not about one country: statistical manuals, the Kyoto Convention, methods papers.",
                "News: no news items have been collected yet. The type is listed so the gap is visible.",
            ],
        },
        "by_country_type": by,
        "totals_by_type": totals,
        "documents": docs,
    }


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True, help="KiriTrade folder")
    ap.add_argument("--out", required=True, help="output JSON path")
    a = ap.parse_args(argv)
    res = build(Path(a.root))
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text(json.dumps(res, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    m = res["meta"]
    print(f"corpus.json: {m['n_documents']} documents, {m['n_words']:,} words -> {a.out}")
    for k in TYPE_ORDER:
        t = res["totals_by_type"][k]
        print(f"  {TYPES[k]:36s} {t['documents']:3d} docs {t['words']:>10,} words")
    return 0


if __name__ == "__main__":
    sys.exit(main())
